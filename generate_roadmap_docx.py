import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette
    COLOR_PRIMARY = RGBColor(30, 58, 138)     # Deep Navy #1E3A8A
    COLOR_ACCENT = RGBColor(14, 116, 144)     # Teal #0E7490
    COLOR_DARK = RGBColor(31, 41, 55)         # Charcoal #1F2937
    COLOR_MUTED = RGBColor(107, 114, 128)     # Gray #6B7280

    # 1. Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("RAASTA KASHMIR")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(4)
    run_sub = sub_p.add_run("Master Technical Blueprint & Backend Architecture Roadmap")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_ACCENT
    run_sub.font.bold = True

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(16)
    run_meta = meta_p.add_run("Internal Engineering Guide for Core Development Team & Interns · Document Version 2.1")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = COLOR_MUTED

    # 2. Executive Overview
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. Executive Overview & Target Ecosystem")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(8)
    r1 = p1.add_run("Raasta Kashmir is an intelligent school-bus safety platform designed for the unique infrastructure and terrain of Kashmir. The final ecosystem consists of 4 distinct components operating in real time:")
    r1.font.name = "Arial"
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_DARK

    bullets = [
        ("Driver Mobile App (Android):", " Dashboard-mounted smartphone running on-device MediaPipe AI vision via an external USB-OTG camera (monitoring fatigue, microsleep, mobile phone distraction, and seatbelt compliance) alongside GPS tracking."),
        ("Parent Mobile App (Android/iOS):", " Real-time parent tracking portal showing the live bus location, stop arrival countdowns (<5-minute proximity alerts), and student drop-off confirmations."),
        ("School Admin & RTO Web Portal (Desktop Browser):", " High-density command center allowing school transport managers and traffic regulators to monitor all active buses simultaneously, review driver safety compliance scores, and audit flagged incident snapshots."),
        ("Central Company Server (Ubuntu Linux):", " A 24/7 cloud server hosting a Node.js API, a sub-second Socket.io WebSocket engine for live GPS broadcast, and a persistent PostgreSQL database.")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        rt = bp.add_run(b_title)
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = COLOR_PRIMARY
        rd = bp.add_run(b_desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(9.5)
        rd.font.color.rgb = COLOR_DARK

    # 3. Master Summary Table
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. The 4-Phase Roadmap Summary")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.1), Inches(1.1), Inches(2.7), Inches(1.8)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    headers = ["Phase", "Timeline", "Core Deliverables", "Infrastructure Used"]
    hdr_row = table.rows[0]
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    phase_data = [
        ("Phase 1", "Immediate\n(Tomorrow)", "Deterministic V1 Demo: multi-role UI, real-world GPS simulation along Dal Lake, live MediaPipe vision AI cockpit, and admin incident forwarding.", "Local Machine / Physical Phone (Zero cloud dependency)"),
        ("Phase 2", "Weeks 1 to 3\n(Backend Focus)", "Company Server & Database: Node.js API, PostgreSQL schema, Socket.io real-time GPS relays (<500ms), and automated DMS violation logging.", "Company Linux Server\n(Ubuntu + Static IP)"),
        ("Phase 3", "Weeks 4 to 6", "Platform Specialization: Standalone Android APK for drivers with background foreground service & offline SQLite queue; desktop web portal for school admins.", "EAS Build APK +\nDesktop Web Dashboard"),
        ("Phase 4", "Weeks 7 to 8", "Vehicle Hardware Deployment: External USB-OTG wide-angle/IR camera, fast-charging power passthrough, Kashmir road bump filtering & live bus pilot.", "Operational School Bus +\nUSB IR Camera")
    ]

    for row_idx, data in enumerate(phase_data, start=1):
        row = table.rows[row_idx]
        bg_color = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.bold = True if col_idx == 0 else False
            r.font.color.rgb = COLOR_PRIMARY if col_idx == 0 else COLOR_DARK

    # 4. Dedicated Backend Section (NEW & DETAILED)
    doc.add_page_break()
    h_be = doc.add_paragraph()
    h_be.paragraph_format.space_before = Pt(10)
    h_be.paragraph_format.space_after = Pt(4)
    r_hbe = h_be.add_run("3. The Backend Architecture: When, Where & Why We Need It")
    r_hbe.font.name = "Arial"
    r_hbe.font.size = Pt(14)
    r_hbe.font.bold = True
    r_hbe.font.color.rgb = COLOR_PRIMARY

    p_be_intro = doc.add_paragraph()
    p_be_intro.paragraph_format.space_after = Pt(8)
    r_bei = p_be_intro.add_run("In software engineering, understanding where a backend is required is fundamental to moving from a local prototype to a multi-user production system. We write the backend in Phase 2 (Weeks 1 to 3). Here is the precise architectural breakdown of where and why the backend is necessary:")
    r_bei.font.name = "Arial"
    r_bei.font.size = Pt(9.5)
    r_bei.font.color.rgb = COLOR_DARK

    be_points = [
        ("1. Real-Time Low-Latency GPS Streaming (Socket.io Engine):", 
         "A moving school bus cannot connect directly to 40 different parents' phones. The driver's phone sends its coordinates {busId, lat, lng, speed} every 2 seconds to the central server. The backend instantly broadcasts this location to subscribed parents in under 500 milliseconds. Normal HTTP requests are too slow for vehicle tracking; WebSockets keep a persistent open line."),
        ("2. Safety Incident Logging & Evidence Archival:", 
         "When the on-device DMS detects a critical safety violation (e.g., driver asleep for >1.5s or holding a phone to their ear), the mobile app cannot store this alone. The backend ingests the incident payload (timestamp, bus ID, AI confidence score, and compressed photo snapshot) and permanently commits it to the database for school administrative audit."),
        ("3. Trip Lifecycle & Route Governance:", 
         "The backend manages active vs. completed trips, calculates official start/end timestamps, detects unauthorized route detours, and records peak speeds to maintain government RTO transport compliance."),
        ("4. Authentication & Role-Based Access Control (RBAC):", 
         "Security requires that parents only see THEIR assigned child's bus and pickup schedule, drivers only see THEIR assigned route, and school admins have district-wide visibility. The backend verifies credentials using JWT tokens with encrypted password hashing (bcrypt)."),
        ("5. Student Roster & Live Attendance Tracking:", 
         "When students board or exit the bus, the driver's app confirms their status. The backend updates the student's live status (HOME -> ON_BUS -> AT_SCHOOL) and triggers push notifications to parents.")
    ]

    for pt_title, pt_desc in be_points:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        rt = p.add_run(pt_title + " ")
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = COLOR_ACCENT
        rd = p.add_run(pt_desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(9.5)
        rd.font.color.rgb = COLOR_DARK

    # 5. Deep Dive into the 4 Phases
    phases = [
        ("4. Phase 1: Frontend Architecture & Local AI Validation (Immediate / Tomorrow)", [
            ("Objective:", "Prove the complete safety workflow, user journeys, and computer vision reliability locally without depending on live cloud infrastructure."),
            ("What is Working in D:\\RAAH:", "The DMS Core (17 specialized modules from Driver-Attention) is integrated in src/services/dms/core/. It computes PERCLOS eye closure, calling vs. texting phone classification, seatbelt hybrid fusion, and steering zone compliance."),
            ("The 5-Minute Demo Flow:", "1) Log in as Driver (driver@raasta.com) and start the Dal Lake route simulation. 2) Tap 'Simulate Detour' to show route diversion alerts. 3) Launch 'AI Cockpit' to test live on-camera Microsleep (close eyes >1.5s) and phone-at-ear detection. 4) Switch to School Admin (admin@raasta.com) to show the violation recorded in the Incident Review Queue with timestamps and AI confidence scores."),
            ("Why Zero Backend Today:", "Prevents network connectivity drops from disrupting the presentation. Proves pure software capability first.")
        ]),
        ("5. Phase 2: Live Company Server Deployment & Telemetry Integration (Weeks 1 to 3)", [
            ("Objective:", "Transform Raasta Kashmir from a single-device simulation into a multi-device live production system communicating over the public internet."),
            ("Backend & Database Stack:", "Node.js (Express), PostgreSQL 15+ (normalized tables for users, buses, routes, route_stops, students, trips, dms_incidents), Redis (geospatial in-memory caching), Nginx (SSL encryption), and PM2 (24/7 uptime)."),
            ("What to Ask Sir After the Demo:", "1) SSH access to the company server (IP, username, credentials/key with root access). 2) Google Cloud Maps API key (with Maps SDK for Android & Maps JavaScript enabled). 3) Subdomain pointing to the server (e.g., api.raasta.company.com).")
        ]),
        ("6. Phase 3: Platform Specialization - Mobile APK vs. Web Portal (Weeks 4 to 6)", [
            ("Objective:", "Split the unified codebase into platform-specialized production applications."),
            ("Driver & Parent Mobile Apps:", "Packaged as standalone Android APKs. Drivers get an Android Foreground Service so GPS and DMS continue running if Google Maps is opened. Includes an offline SQLite queue so coordinates are saved during mountain pass cellular blackouts and auto-synced upon 4G recovery."),
            ("School Admin & RTO Web Portal:", "Desktop-optimized browser application displaying all fleet buses on a multi-bus vector map with live color-coded compliance status (Green: Normal, Yellow: Detour, Red: DMS Violation).")
        ]),
        ("7. Phase 4: Physical Bus Hardware Setup & Field Road Pilot (Weeks 7 to 8)", [
            ("Objective:", "Mount real hardware in an operational school bus in Srinagar and conduct live field testing."),
            ("Hardware Installation:", "Mount an external wide-angle (110°-130° FOV) 850nm Infrared (IR) night-vision USB webcam on the dashboard. Connect it to the driver's phone via a USB-C Power Delivery (PD) OTG splitter, fast-charging the phone from the bus 12V socket while streaming 640x480 video at 15 FPS."),
            ("Zero Overheating Architecture:", "Uses Android native UVC drivers (UsbCameraDmsModule.kt) and native MediaPipe tasks, running on the phone's native NPU/GPU to prevent thermal throttling."),
            ("Steering Wheel & Hands-On Calibration System:", "Hands-on-wheel detection cannot be calibrated during Phase 1-3 because camera mounting height, cabin geometry, and angle of view vary across school bus models (e.g., Tata Winger vs. Force Traveller vs. SML Isuzu). During Phase 4 in-bus deployment, we deploy an in-app 'Steering Calibration Wizard': 1) The driver places hands at the standard 10-and-2 position, 2) MediaPipe Hand Landmark detection records the hand wrist/knuckle coordinate cluster, 3) The bounding 'Hands-On Steering Zone' polygon is dynamically calibrated and saved to the vehicle's profile."),
            ("Kashmir Terrain Tuning:", "Calibrates head pose math to filter out road bumps and potholes so vertical shocks do not cause false microsleep alarms.")
        ])
    ]

    for title, points in phases:
        ph = doc.add_paragraph()
        ph.paragraph_format.space_before = Pt(12)
        ph.paragraph_format.space_after = Pt(4)
        r_ph = ph.add_run(title)
        r_ph.font.name = "Arial"
        r_ph.font.size = Pt(13)
        r_ph.font.bold = True
        r_ph.font.color.rgb = COLOR_PRIMARY

        for pt_title, pt_desc in points:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            rt = p.add_run(pt_title + " ")
            rt.font.name = "Arial"
            rt.font.size = Pt(9.5)
            rt.font.bold = True
            rt.font.color.rgb = COLOR_ACCENT
            rd = p.add_run(pt_desc)
            rd.font.name = "Arial"
            rd.font.size = Pt(9.5)
            rd.font.color.rgb = COLOR_DARK

    out_path = r"D:\RAAH\RAASTA_4_PHASE_ROADMAP_MASTER.docx"
    doc.save(out_path)
    print("Master DOCX updated successfully with backend specifications at:", out_path)

if __name__ == "__main__":
    create_document()
